"""Pull a small text snippet out of a document, locally.

Only enough context for the model to tell a SAT paper from a dissertation
chapter. Content is always treated as data: see classify.py for the wrapping.
"""
from __future__ import annotations

import re
import logging
import unicodedata
import zipfile
from pathlib import Path

from . import host
from .scanner import FileRecord

# Two guards, because a single pathological file must not stall a daily run:
#   * a size ceiling — a 111 MB scanned textbook has no useful first page and
#     costs minutes to parse
#   * a per-file wall clock — pypdf can crawl on a damaged or huge page tree
logging.getLogger("pypdf").setLevel(logging.ERROR)

MAX_EXTRACT_BYTES = 25 * 1024 * 1024
EXTRACT_TIMEOUT_SECONDS = 20


TEXTLIKE = {"txt", "md", "markdown", "csv", "tsv", "json", "yaml", "yml",
            "html", "htm", "xml", "rtf", "log", "tex", "bib", "srt", "vtt"}
IMAGE = {"png", "jpg", "jpeg", "gif", "heic", "webp", "tiff", "bmp", "svg"}
VIDEO = {"mp4", "mov", "avi", "mkv", "webm", "m4v", "mpg", "mpeg", "wmv"}
AUDIO = {"mp3", "m4a", "wav", "aac", "flac", "aiff", "ogg", "opus", "wma"}
ARCHIVE = {"zip", "tar", "gz", "tgz", "bz2", "xz", "7z", "rar"}
INSTALLER = {"dmg", "pkg", "exe", "msi", "deb", "rpm", "appimage"}

_WS = re.compile(r"[ \t ]+")
_NL = re.compile(r"\n{3,}")
_TAG = re.compile(r"<[^>]{1,400}>")


def kind_of(ext: str) -> str:
    if ext in IMAGE:
        return "image"
    if ext in VIDEO:
        return "video"
    if ext in AUDIO:
        return "audio"
    if ext in ARCHIVE:
        return "archive"
    if ext in INSTALLER:
        return "installer"
    if ext == "pdf":
        return "pdf"
    if ext in {"doc", "docx", "odt", "pages"}:
        return "document"
    if ext in {"xls", "xlsx", "ods", "numbers"}:
        return "spreadsheet"
    if ext in {"ppt", "pptx", "odp", "key"}:
        return "presentation"
    if ext in TEXTLIKE:
        return "text"
    return "other"


def _clean(text: str, limit: int) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = "".join(ch for ch in text if ch == "\n" or ch == "\t" or ch.isprintable())
    text = _WS.sub(" ", text)
    text = _NL.sub("\n\n", text).strip()
    return text[:limit]


def _pdf(path: Path, limit: int) -> tuple[str, str]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "", "pypdf unavailable"
    try:
        reader = PdfReader(str(path), strict=False)
        n = len(reader.pages)
        parts: list[str] = []
        for page in reader.pages[:4]:
            parts.append(page.extract_text() or "")
            if sum(len(p) for p in parts) >= limit:
                break
        text = _clean("\n".join(parts), limit)
        note = f"{n} pages"
        if not text:
            note += "; no extractable text (likely a scan)"
        return text, note
    except Exception as e:
        return "", f"unreadable pdf: {type(e).__name__}"


def _docx(path: Path, limit: int) -> tuple[str, str]:
    try:
        import docx
    except ImportError:
        return "", "python-docx unavailable"
    try:
        d = docx.Document(str(path))
        parts, total = [], 0
        for p in d.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            parts.append(t)
            total += len(t)
            if total >= limit:
                break
        return _clean("\n".join(parts), limit), f"{len(d.paragraphs)} paragraphs"
    except Exception as e:
        return "", f"unreadable docx: {type(e).__name__}"


def _legacy_doc(path: Path, limit: int) -> tuple[str, str]:
    """Best-effort for .doc: pull readable runs out of the binary."""
    try:
        raw = path.open("rb").read(400_000)
    except OSError as e:
        return "", f"unreadable: {e.strerror}"
    runs = re.findall(rb"[\x20-\x7e]{6,}", raw)
    text = " ".join(r.decode("ascii", "ignore") for r in runs[:400])
    return _clean(text, limit), "legacy .doc, extracted heuristically"


def _plain(path: Path, limit: int, strip_tags: bool = False) -> tuple[str, str]:
    try:
        raw = path.open("r", encoding="utf-8", errors="replace").read(limit * 6)
    except OSError as e:
        return "", f"unreadable: {e.strerror}"
    if strip_tags:
        raw = _TAG.sub(" ", raw)
    return _clean(raw, limit), ""


def _zip_listing(path: Path, limit: int) -> tuple[str, str]:
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
        head = ", ".join(names[:25])
        return _clean(head, limit), f"{len(names)} entries"
    except Exception as e:
        return "", f"archive not listable: {type(e).__name__}"


def _dispatch(record: FileRecord, limit: int) -> tuple[str, str]:
    ext, kind = record.ext, record.extra["kind"]
    if kind == "pdf":
        return _pdf(record.path, limit)
    if ext == "docx":
        return _docx(record.path, limit)
    if ext == "doc":
        return _legacy_doc(record.path, limit)
    if ext in {"html", "htm", "xml"}:
        return _plain(record.path, limit, strip_tags=True)
    if ext in TEXTLIKE:
        return _plain(record.path, limit)
    if ext in ARCHIVE or ext in {"zip", "xlsx", "pptx"}:
        return _zip_listing(record.path, limit)
    return "", ""


def enrich(record: FileRecord, limit: int) -> None:
    """Attach a text snippet and a short note to a record, in place.

    Failure is never fatal: the model can still classify from the filename,
    type and size alone, which for most files is most of the signal anyway.
    """
    kind = kind_of(record.ext)
    record.extra["kind"] = kind

    if record.size > MAX_EXTRACT_BYTES and kind not in ("image", "video", "audio"):
        record.snippet = ""
        record.snippet_note = (f"{record.size / 1048576:.0f}MB, too large to read "
                               "— classified on name and type alone")
        return

    try:
        record.snippet, record.snippet_note = host.run_with_timeout(
            lambda: _dispatch(record, limit), EXTRACT_TIMEOUT_SECONDS)
    except TimeoutError:
        record.snippet = ""
        record.snippet_note = (f"no text within {EXTRACT_TIMEOUT_SECONDS}s "
                               "— classified on name and type alone")
    except Exception as e:
        record.snippet = ""
        record.snippet_note = f"extraction failed: {type(e).__name__}"
